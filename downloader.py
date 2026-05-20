"""
downloader.py — Descarga los 8 reportes de Retail Analytics
"""

import time, random, logging, os, re, json
from datetime import date, timedelta
from pathlib import Path
from playwright.sync_api import Page

from config import (
    OUTPUT_DIR, POLL_INTERVAL_SEC, REPORT_TIMEOUT_SEC, DELAY_BETWEEN_REQUESTS
)

logger = logging.getLogger(__name__)
BASE = "https://vendorcentral.amazon.com/retail-analytics/dashboard"

REPORT_URLS = {
    "sales":        f"{BASE}/sales",
    "realtime":     f"{BASE}/real-time-sales-v2",
    "inventory":    f"{BASE}/inventory",
    "traffic":      f"{BASE}/traffic",
    "forecasting":  f"{BASE}/forecasting",
    "df_forecast":  f"{BASE}/df-forecasting",
    "netppm":       f"{BASE}/net-ppm",
    "catalog":      f"{BASE}/product-catalog",
}


def _make_retail_url(report_key: str, account: dict = None) -> str:
    """Devuelve la URL del reporte, añadiendo entityId si la cuenta lo tiene configurado."""
    url = REPORT_URLS[report_key]
    entity_id = (account or {}).get("retail_entity_id")
    if entity_id:
        sep = "&" if "?" in url else "?"
        return f"{url}{sep}entityId={entity_id}"
    return url


_JS_GET_ALL = """function getAll(root) {
    var nodes = [];
    try {
        var tw = document.createTreeWalker(root, NodeFilter.SHOW_ELEMENT);
        var n; while (n = tw.nextNode()) {
            nodes.push(n);
            if (n.shadowRoot) nodes = nodes.concat(getAll(n.shadowRoot));
        }
    } catch(e) {}
    return nodes;
}"""


def _read_current_vc_account(page) -> str:
    """Lee el nombre de cuenta activo en el header de Vendor Central."""
    try:
        return page.evaluate(f"""() => {{
            {_JS_GET_ALL}
            var all = getAll(document);
            var best = '', bestLen = 9999;
            for (var i = 0; i < all.length; i++) {{
                var t = (all[i].textContent || '').trim();
                if (t.startsWith('US - ') && t.length < 100 && t.length < bestLen) {{
                    best = t; bestLen = t.length;
                }}
            }}
            return best;
        }}""")
    except Exception:
        return ""


def switch_vendor_central_account(page: Page, account: dict) -> bool:
    """
    Cambia a la cuenta indicada en Vendor Central vía el switcher del header.
    Paso 1: click en el header (y < 80) para abrir la lista de cuentas.
    Paso 2: click en la cuenta destino en la lista (y > 60, matching flexible).
    """
    header_name = (account or {}).get("vc_header_name", "")
    if not header_name:
        return True
    account_short = header_name.replace("US - ", "")

    # ── Verificar cuenta actual ───────────────────────────────────────────────
    current = _read_current_vc_account(page)
    if current == header_name:
        logger.info(f"✓ Vendor Central ya en cuenta: {header_name}")
        return True
    if current:
        logger.info(f"Cuenta actual: '{current}' → cambiando a '{header_name}'")

    # ── Paso 1: click en el botón del header ─────────────────────────────────
    # Preferir el elemento con la y MÁS ALTA (más específico/hoja del DOM),
    # y entre iguales preferir SPAN/BUTTON/A sobre DIV — evita clickear
    # el contenedor padre en lugar del botón real.
    click_text = current if current else header_name
    result_step1 = page.evaluate(f"""([txt]) => {{
        {_JS_GET_ALL}
        var all = getAll(document);
        var best = null, bestTop = -1, bestPri = 0;
        var TAG_PRI = {{SPAN:3, BUTTON:3, A:3, DIV:1}};
        var log = [];
        for (var i = 0; i < all.length; i++) {{
            var t = (all[i].textContent || '').trim();
            if (t === txt) {{
                var r = all[i].getBoundingClientRect();
                log.push(r.top.toFixed(0) + '[' + all[i].tagName + ']');
                if (r.top >= 0 && r.top < 80 && r.height > 0) {{
                    var pri = TAG_PRI[all[i].tagName] || 1;
                    if (r.top > bestTop || (r.top === bestTop && pri > bestPri)) {{
                        bestTop = r.top; bestPri = pri; best = all[i];
                    }}
                }}
            }}
        }}
        if (best) {{
            best.click();
            return 'clicked y=' + bestTop.toFixed(0) + ' [' + best.tagName + ']';
        }}
        return 'not_found|' + log.slice(0,8).join(';');
    }}""", [click_text])
    logger.info(f"Paso1 click header: {result_step1}")

    if result_step1.startswith("not_found"):
        logger.warning("No se encontró el botón de cuenta en el header.")
        page.screenshot(path="debug_vc_no_trigger.png")
        return False

    # Esperar a que el SPA cargue la lista de cuentas (NO usar networkidle — tarda 40s)
    time.sleep(4)
    page.screenshot(path="debug_vc_switch_accounts.png")

    # ── Paso 2: seleccionar la cuenta destino ─────────────────────────────────
    for attempt in range(12):
        # Intento A: Playwright get_by_text (auto-espera, pierces shadow DOM)
        for name_try in [header_name, account_short]:
            try:
                page.get_by_text(name_try, exact=True).first.click(timeout=2000)
                logger.info(f"✓ Cuenta seleccionada via get_by_text('{name_try}')")
                try:
                    page.wait_for_load_state("networkidle", timeout=30000)
                except Exception:
                    time.sleep(6)
                new_account = _read_current_vc_account(page)
                logger.info(f"Cuenta tras switch: '{new_account}'")
                return True
            except Exception:
                pass

        # Intento B: JS getBoundingClientRect (fallback shadow DOM profundo)
        # y > 60 evita re-clickear el header (double-click cierra el panel)
        result_step2 = page.evaluate(f"""([full, short]) => {{
            {_JS_GET_ALL}
            var all = getAll(document);
            for (var i = 0; i < all.length; i++) {{
                var t = (all[i].textContent || '').replace(/\\s+/g, ' ').trim();
                if ((t === full || t === short || t.includes(full) ||
                     (t.startsWith(short) && t.length < 120)) && t.length < 200) {{
                    var r = all[i].getBoundingClientRect();
                    if (r.top > 60 && r.height > 0) {{
                        all[i].click();
                        return 'js_clicked:' + t.substring(0,60) + ' y=' + r.top.toFixed(0) + ' [' + all[i].tagName + ']';
                    }}
                }}
            }}
            var cands = [];
            for (var i = 0; i < all.length; i++) {{
                var t = (all[i].textContent || '').replace(/\\s+/g, ' ').trim();
                if ((t.includes('Can Do') || t.includes('Honey') || t.includes('US - ')) && t.length < 120) {{
                    var r = all[i].getBoundingClientRect();
                    cands.push('"' + t.substring(0,35) + '" y=' + r.top.toFixed(0) + '[' + all[i].tagName + ']');
                }}
            }}
            return 'not_found|' + cands.slice(0,10).join(' | ');
        }}""", [header_name, account_short])

        if not result_step2.startswith("not_found"):
            logger.info(f"✓ Cuenta seleccionada: {result_step2}")
            try:
                page.wait_for_load_state("networkidle", timeout=30000)
            except Exception:
                time.sleep(6)
            new_account = _read_current_vc_account(page)
            logger.info(f"Cuenta tras switch: '{new_account}'")
            return True

        logger.info(f"Intento {attempt + 1}/12: {result_step2}")
        time.sleep(2)

    logger.warning(f"No se pudo seleccionar '{header_name}' después de 12 intentos.")
    page.screenshot(path="debug_vc_option_not_found.png")
    return False

REPORT_H1_TEXT = {
    "sales":     "Sales",
    "inventory": "Inventory",
    "traffic":   "Traffic",
    "netppm":    "Net PPM",
}

# JavaScript helper para recorrer shadow DOM (también lo importa ads_downloader.py)
_GET_ALL_FN = """function getAll(root) {
    var nodes = [];
    try {
        var tw = document.createTreeWalker(root, NodeFilter.SHOW_ELEMENT);
        var n; while (n = tw.nextNode()) { nodes.push(n); if (n.shadowRoot) nodes = nodes.concat(getAll(n.shadowRoot)); }
    } catch(e) {}
    return nodes;
}"""


# ── Utilidades de página ──────────────────────────────────────────────────────

def ensure_page(context, page):
    """Devuelve la página si sigue activa; si no, abre una nueva."""
    try:
        page.evaluate("1")
        return page
    except Exception:
        return context.new_page()


def upload_file(filepath: str, account: dict = None):
    """Sube el archivo a SharePoint en la carpeta correcta para la cuenta indicada."""
    try:
        from uploader import upload_to_sharepoint
        upload_to_sharepoint(filepath, account)
    except Exception as e:
        logger.warning(f"upload_file: {e}")


def stamp_download_date(filepath: str):
    """
    Inyecta la fecha/hora de descarga en el archivo Excel:
      - Propiedad del documento (Description)
      - Hoja 'Metadata' con fecha y hora legibles
    """
    try:
        import openpyxl
        from datetime import datetime
        wb = openpyxl.load_workbook(filepath)
        now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        wb.properties.description = f"Download Date: {now}"
        if "Metadata" not in wb.sheetnames:
            ws = wb.create_sheet("Metadata")
        else:
            ws = wb["Metadata"]
        ws["A1"] = "Download Date"
        ws["B1"] = now
        wb.save(filepath)
        logger.info(f"✓ Fecha de descarga añadida: {now}")
    except ImportError:
        logger.warning("openpyxl no instalado — fecha no añadida (pip install openpyxl)")
    except Exception as e:
        logger.warning(f"stamp_download_date: {e}")


def random_delay():
    time.sleep(random.uniform(*DELAY_BETWEEN_REQUESTS))

def setup_output_dir():
    Path(OUTPUT_DIR).mkdir(parents=True, exist_ok=True)


# ── KAT dropdowns ─────────────────────────────────────────────────────────────

def js_open_dropdown(page: Page, current_value: str):
    page.evaluate(f"""() => {{
        function getAll(root) {{
            let all = [];
            const w = document.createTreeWalker(root, NodeFilter.SHOW_ELEMENT);
            let n; while (n = w.nextNode()) {{ all.push(n); if (n.shadowRoot) all = all.concat(getAll(n.shadowRoot)); }}
            return all;
        }}
        for (const el of getAll(document)) {{
            if (el.tagName === 'SLOT' && el.name === 'selected-option' && el.textContent.includes('{current_value}')) {{
                let n = el;
                while (n) {{ if (n.tagName && n.tagName.startsWith('KAT-')) {{ n.click(); return; }} n = n.parentElement || (n.getRootNode && n.getRootNode().host); }}
            }}
        }}
    }}""")
    time.sleep(1.5)

def js_click_option(page: Page, text: str) -> bool:
    result = page.evaluate(f"""() => {{
        function getAll(root) {{
            let all = [];
            const w = document.createTreeWalker(root, NodeFilter.SHOW_ELEMENT);
            let n; while (n = w.nextNode()) {{ all.push(n); if (n.shadowRoot) all = all.concat(getAll(n.shadowRoot)); }}
            return all;
        }}
        for (const el of getAll(document)) {{
            if (typeof el.className === 'string' && el.className.includes('standard-option-name') && el.textContent.trim() === '{text}') {{
                el.click(); return 'clicked';
            }}
        }}
        return 'not found';
    }}""")
    time.sleep(0.8)
    return result == 'clicked'

def kat_select(page: Page, current: str, desired: str) -> bool:
    logger.info(f"Dropdown: '{current}' → '{desired}'")
    js_open_dropdown(page, current)
    ok = js_click_option(page, desired)
    if not ok:
        time.sleep(0.5)
        ok = js_click_option(page, desired)
    return ok


# ── Fecha ──────────────────────────────────────────────────────────────────────

def set_date(page: Page, target_date: date):
    date_str = target_date.strftime("%m/%d/%Y")
    time.sleep(1)
    try:
        inputs = page.locator('input[type="text"]:visible').all()
        for inp in inputs:
            ph = (inp.get_attribute('placeholder') or '').lower()
            if 'search' in ph or 'asin' in ph:
                continue
            val = inp.input_value()
            inp.click(click_count=3)
            time.sleep(0.3)
            inp.type(date_str)
            time.sleep(0.3)
            inp.press('Escape')
            time.sleep(0.3)
            inp.press('Tab')
            time.sleep(0.5)
            logger.info(f"Fecha: {date_str} (era: '{val}')")
            return
        logger.warning("set_date: no input encontrado")
    except Exception as e:
        logger.warning(f"set_date error: {e}")


# ── Apply ──────────────────────────────────────────────────────────────────────

def click_apply(page: Page):
    try:
        apply = page.locator('button:has-text("Apply")').first
        apply.wait_for(timeout=5000, state="visible")
        for _ in range(15):
            if apply.is_enabled():
                break
            time.sleep(1)
        apply.click()
        logger.info("Apply clickeado.")
        time.sleep(3)
    except Exception as e:
        logger.warning(f"Apply error: {e}")


# ── Customize Columns ─────────────────────────────────────────────────────────

def columns_are_complete(page: Page) -> bool:
    col = page.locator('text=/Displaying \\d+ of \\d+ columns/i').first
    if col.count() == 0:
        return False
    txt = col.inner_text()
    nums = re.findall(r'\d+', txt)
    ok = len(nums) == 2 and nums[0] == nums[1]
    logger.info(f"Columnas: {txt} → {'✓' if ok else '✗'}")
    return ok

def select_all_columns(page: Page):
    logger.info("Customize Columns → Select All → Save...")
    page.evaluate("""() => {
        const all = document.querySelectorAll('*');
        for (const el of all) {
            if (el.textContent && el.textContent.trim().startsWith('Customize Columns')) {
                el.click(); return;
            }
        }
    }""")
    time.sleep(2)
    try:
        cb = page.locator('kat-checkbox[label="Select All"]').first
        cb.wait_for(timeout=5000, state="visible")
        cb.click()
        logger.info("Select All clickeado.")
        time.sleep(0.8)
    except Exception as e:
        logger.warning(f"Select All error: {e}")
    page.locator('button:has-text("Save")').first.wait_for(timeout=5000, state="visible")
    page.locator('button:has-text("Save")').first.click()
    time.sleep(2)
    logger.info("✓ Columnas guardadas.")


# ── Excel + Download ───────────────────────────────────────────────────────────

def click_excel(page: Page) -> bool:
    try:
        excel = page.locator('button:has-text("Excel")').first
        excel.wait_for(timeout=8000, state="visible")
        excel.click()

        time.sleep(2)

        warn = page.locator('text=/exceeded/i, text=/superado/i').first

        if warn.count() > 0:
            logger.warning("Límite detectado — esperando 6 min...")
            time.sleep(360)

            excel.click()
            time.sleep(2)

        logger.info("✓ Excel solicitado.")
        return True

    except Exception as e:
        logger.error(f"Excel error: {e}")
        return False


# JavaScript para encontrar el Download link EXACTO de un reporte en el panel.
# Evita el bug de ancestor::div[4] que devuelve el panel entero y matchea cualquier fila.
# Estrategia: encuentra un elemento PEQUEÑO con el texto buscado (una celda, no el panel),
# luego sube por el DOM hasta encontrar el ancestro MÁS PEQUEÑO que tenga ESE texto
# Y un link "Download" — eso es la fila correcta.
_JS_FIND_AND_CLICK_DOWNLOAD = """(searchTerm) => {
    function getAll(root) {
        var nodes = [];
        try {
            var tw = document.createTreeWalker(root, NodeFilter.SHOW_ELEMENT);
            var n; while (n = tw.nextNode()) {
                nodes.push(n);
                if (n.shadowRoot) nodes = nodes.concat(getAll(n.shadowRoot));
            }
        } catch(e) {}
        return nodes;
    }
    function crossParent(el) {
        if (el.parentElement) return el.parentElement;
        var r = el.getRootNode();
        return (r && r !== document && r.host) ? r.host : null;
    }

    var all = getAll(document);
    var term = searchTerm.toLowerCase();

    // Paso 1: encontrar elementos PEQUEÑOS que contengan nuestro término
    // (celdas individuales de la fila, no contenedores grandes del panel)
    var candidates = [];
    for (var i = 0; i < all.length; i++) {
        var el = all[i];
        var txt = (el.textContent || '').trim();
        if (!txt.toLowerCase().includes(term)) continue;
        if (txt.length > 200) continue;         // descartar contenedores grandes
        if (el.childElementCount > 4) continue; // descartar elementos complejos
        candidates.push(el);
    }

    if (candidates.length === 0) {
        var inProg = all.filter(function(n) {
            var t = (n.textContent||'').trim().toLowerCase();
            return t==='in progress'||t==='generating'||t==='processing'||t==='pending';
        }).length;
        return JSON.stringify({status:'term_not_found', term:searchTerm, in_progress:inProg});
    }

    // Paso 2: para cada candidato, subir por el DOM buscando el ancestro MÁS PEQUEÑO
    // que tenga tanto nuestro término como un link "Download"
    for (var ci = 0; ci < candidates.length; ci++) {
        var ancestor = crossParent(candidates[ci]);
        for (var up = 0; up < 12 && ancestor; up++) {
            var aTxt = (ancestor.textContent || '').trim();

            // Si este ancestro no tiene nuestro término, subir más
            if (!aTxt.toLowerCase().includes(term)) {
                ancestor = crossParent(ancestor);
                continue;
            }
            // Si el ancestro es demasiado grande, es el panel completo — subir más
            if (aTxt.length > 1200) {
                ancestor = crossParent(ancestor);
                continue;
            }

            // Buscar link "Download" dentro de este ancestro
            var sub = getAll(ancestor);
            for (var si = 0; si < sub.length; si++) {
                var node = sub[si];
                var nTxt = (node.textContent || '').trim();
                var nTag = (node.tagName || '').toLowerCase();
                if (nTxt === 'Download' && (nTag==='a'||nTag==='button'||nTag==='span')) {
                    node.click();
                    return JSON.stringify({
                        status: 'clicked',
                        term: searchTerm,
                        match: candidates[ci].textContent.trim().substring(0,80),
                        rowText: aTxt.substring(0,150)
                    });
                }
            }
            ancestor = crossParent(ancestor);
        }
    }

    var inProgress = all.filter(function(n) {
        var t = (n.textContent||'').trim().toLowerCase();
        return t==='in progress'||t==='generating'||t==='processing'||t==='pending';
    }).length;
    return JSON.stringify({status:'no_download_link', term:searchTerm, in_progress:inProgress});
}"""


def download_from_panel(page: Page, report_name: str, filename: str = None,
                        output_dir: str = None) -> tuple:
    """
    Busca en Manage Downloads el reporte EXACTO por nombre y lo descarga.
    Usa shadow DOM traversal para identificar la fila correcta
    (evita el bug de ancestor::div[4] que capturaba el panel completo).
    Retorna (success: bool, filepath: str | None).
    """
    output_dir = output_dir or OUTPUT_DIR
    Path(output_dir).mkdir(parents=True, exist_ok=True)
    deadline = time.time() + REPORT_TIMEOUT_SEC
    logger.info(f"Esperando reporte '{report_name}' en Manage Downloads...")

    while time.time() < deadline:
        # Abrir el panel de downloads
        try:
            page.evaluate("document.getElementById('downloadManager').click()")
            time.sleep(2)
        except Exception as e:
            logger.warning(f"No pude abrir panel: {e}")
            time.sleep(5)
            continue

        # Buscar la fila exacta del reporte y descargar
        result_raw = None
        try:
            with page.expect_download(timeout=15000) as dl_info:
                result_raw = page.evaluate(_JS_FIND_AND_CLICK_DOWNLOAD, report_name)

            # Si llegamos aquí sin excepción, se inició una descarga
            try:
                logger.info(f"JS resultado: {json.loads(result_raw)}")
            except Exception:
                pass

            dl = dl_info.value
            amazon_filename = dl.suggested_filename
            save_name = filename if filename else amazon_filename
            filepath = os.path.join(output_dir, save_name)
            logger.info(f"Nombre Amazon: {amazon_filename} → guardando como: {save_name}")
            dl.save_as(filepath)
            logger.info(f"✓ Guardado: {filepath}")
            return True, filepath

        except Exception as e:
            err_str = str(e).lower()
            if "timeout" in err_str:
                if result_raw:
                    try:
                        logger.info(f"Estado panel: {json.loads(result_raw)}")
                    except Exception:
                        logger.info(f"Estado panel: {result_raw}")
                else:
                    logger.info(f"Reporte '{report_name}' aún no disponible.")
            else:
                logger.warning(f"download_from_panel: {e}")

        logger.info(f"No listo. Esperando {POLL_INTERVAL_SEC}s...")
        try:
            page.keyboard.press("Escape")
        except Exception:
            pass
        time.sleep(POLL_INTERVAL_SEC)

    logger.error(f"Timeout para '{report_name}'")
    return False, None


# ── Reportes con fecha (loop diario) ──────────────────────────────────────────

def download_daily_report(page: Page, report_key: str, target_date: date,
                           extra_filters: dict = None,
                           account: dict = None) -> bool:
    """
    Descarga un reporte diario (Sales, Inventory, Traffic, Net PPM).
    Usa el nombre original de Amazon (ej: Sales_ASIN_Sourcing_Retail_UnitedStates_Daily_1-1-2026_1-1-2026.xlsx).
    """
    url        = _make_retail_url(report_key, account)
    date_tag   = f"{target_date.month}-{target_date.day}-{target_date.year}"
    output_dir = (account or {}).get("output_dir", OUTPUT_DIR)

    existing = list(Path(output_dir).glob(f"*_{date_tag}_{date_tag}.*"))
    if existing:
        logger.info(f"Ya existe: {existing[0].name}")
        return True

    logger.info(f"\n{'='*50}")
    logger.info(f"Reporte: {report_key.upper()} | Fecha: {target_date}")

    page.goto(url, timeout=40000)
    h1 = REPORT_H1_TEXT.get(report_key, report_key.replace("_", " ").title())
    page.locator(f'h1:has-text("{h1}")').first.wait_for(timeout=20000)
    time.sleep(2)

    kat_select(page, "Custom", "Daily")
    time.sleep(1)
    set_date(page, target_date)

    if extra_filters:
        for current, desired in extra_filters.items():
            kat_select(page, current, desired)

    click_apply(page)

    if not columns_are_complete(page):
        select_all_columns(page)
        kat_select(page, "Custom", "Daily")
        time.sleep(1)
        set_date(page, target_date)
        if extra_filters:
            for current, desired in extra_filters.items():
                kat_select(page, current, desired)
        click_apply(page)

    if not click_excel(page):
        return False

    ok, filepath = download_from_panel(page, date_tag, output_dir=output_dir)
    if ok and filepath:
        stamp_download_date(filepath)
        upload_file(filepath, account)
    return ok


def download_static_report(page: Page, report_key: str,
                            filters: dict = None, h1_text: str = None,
                            account: dict = None) -> bool:
    """
    Descarga un reporte sin fecha (Real Time Sales, Forecasting, DF Forecast, Catalog).
    Usa el nombre original de Amazon para el archivo.
    """
    from datetime import datetime
    today    = datetime.now().date()
    date_tag = f"{today.month}/{today.day}/{today.year}"   # "M/D/YYYY" — como Amazon lo muestra en panel

    logger.info(f"\n{'='*50}")
    logger.info(f"Reporte estático: {report_key.upper()} | Buscando en panel: '{date_tag}'")

    url = _make_retail_url(report_key, account)
    page.goto(url, timeout=40000)

    h1 = h1_text or report_key.replace("_", " ").title()
    try:
        page.locator(f'h1:has-text("{h1}")').first.wait_for(timeout=20000)
    except Exception:
        time.sleep(3)
    time.sleep(2)

    if filters:
        for current, desired in filters.items():
            kat_select(page, current, desired)

    click_apply(page)
    time.sleep(2)

    if not columns_are_complete(page):
        select_all_columns(page)
        if filters:
            for current, desired in filters.items():
                kat_select(page, current, desired)
        click_apply(page)

    if not click_excel(page):
        return False

    output_dir = (account or {}).get("output_dir", OUTPUT_DIR)
    ok, filepath = download_from_panel(page, date_tag, output_dir=output_dir)
    if ok and filepath:
        stamp_download_date(filepath)
        upload_file(filepath, account)
    return ok


# ── Orquestador principal ──────────────────────────────────────────────────────

def download_all_reports(page: Page, start: date, end: date,
                          context=None, account: dict = None,
                          reports_to_run=None) -> Page:
    """
    Descarga todos los reportes de Retail Analytics seleccionados.

    reports_to_run: "all"  →  descarga los 8 reportes
                   lista   →  solo los reportes cuya clave esté en la lista
                              ej: ["sales", "inventory", "traffic"]
    """
    output_dir = (account or {}).get("output_dir", OUTPUT_DIR)
    Path(output_dir).mkdir(parents=True, exist_ok=True)

    # Cambiar a la cuenta correcta en Vendor Central antes de descargar
    if account:
        page.goto("https://vendorcentral.amazon.com", timeout=30000)
        page.wait_for_load_state("domcontentloaded", timeout=15000)
        # Esperar a que el botón de cuenta sea visible (shadow DOM traversal)
        try:
            page.wait_for_function(
                """() => {
                    function getAll(root) {
                        var nodes = [];
                        try {
                            var tw = document.createTreeWalker(root, NodeFilter.SHOW_ELEMENT);
                            var n; while (n = tw.nextNode()) {
                                nodes.push(n);
                                if (n.shadowRoot) nodes = nodes.concat(getAll(n.shadowRoot));
                            }
                        } catch(e) {}
                        return nodes;
                    }
                    return getAll(document).some(function(el) {
                        return (el.textContent || '').trim().startsWith('US - ');
                    });
                }""",
                timeout=15000,
            )
        except Exception:
            time.sleep(5)
        switch_vendor_central_account(page, account)

    run_all = (reports_to_run is None or reports_to_run == "all")
    selected = set() if run_all else set(reports_to_run)

    def should_run(key: str) -> bool:
        return run_all or key in selected

    all_dates = []
    d = start
    while d <= end:
        all_dates.append(d)
        d += timedelta(days=1)

    logger.info(f"Total días: {len(all_dates)} ({start} → {end})")
    if not run_all:
        logger.info(f"Reportes seleccionados: {sorted(selected)}")

    total = 0

    # ── 1. Sales ──────────────────────────────────────────────────────────────
    if should_run("sales"):
        logger.info("\n" + "█"*50)
        logger.info("SECCIÓN 1: SALES")
        for d in all_dates:
            ok = download_daily_report(page, "sales", d,
                 extra_filters={"Manufacturing": "Sourcing"}, account=account)
            if ok: total += 1
            random_delay()

    # ── 2. Real Time Sales ────────────────────────────────────────────────────
    if should_run("realtime"):
        logger.info("\n" + "█"*50)
        logger.info("SECCIÓN 2: REAL TIME SALES")
        ok = download_static_report(page, "realtime",
             filters={"Trailing 24 hours": "Trailing 48 hours"},
             h1_text="Real Time Sales", account=account)
        if ok: total += 1

    # ── 3. Inventory ──────────────────────────────────────────────────────────
    if should_run("inventory"):
        logger.info("\n" + "█"*50)
        logger.info("SECCIÓN 3: INVENTORY")
        for d in all_dates:
            ok = download_daily_report(page, "inventory", d,
                 extra_filters={"Manufacturing": "Sourcing"}, account=account)
            if ok: total += 1
            random_delay()

    # ── 4. Traffic ────────────────────────────────────────────────────────────
    if should_run("traffic"):
        logger.info("\n" + "█"*50)
        logger.info("SECCIÓN 4: TRAFFIC")
        for d in all_dates:
            ok = download_daily_report(page, "traffic", d, account=account)
            if ok: total += 1
            random_delay()

    # ── 5. Forecasting ────────────────────────────────────────────────────────
    if should_run("forecasting"):
        logger.info("\n" + "█"*50)
        logger.info("SECCIÓN 5: FORECASTING")
        ok = download_static_report(page, "forecasting",
             filters={"Manufacturing": "Retail", "Mean Forecast": "Mean Forecast"},
             h1_text="Forecasting", account=account)
        if ok: total += 1

    # ── 6. Direct Fulfillment Forecasting ─────────────────────────────────────
    if should_run("df_forecast"):
        logger.info("\n" + "█"*50)
        logger.info("SECCIÓN 6: DIRECT FULFILLMENT FORECASTING")
        has_df = (account or {}).get("has_df_forecast", False)
        if has_df:
            ok = download_static_report(page, "df_forecast",
                 h1_text="Direct Fulfillment Forecasting", account=account)
            if ok: total += 1
        else:
            logger.info("→ Esta cuenta no tiene Direct Fulfillment Forecasting. Saltando.")

    # ── 7. Net PPM ────────────────────────────────────────────────────────────
    if should_run("netppm"):
        logger.info("\n" + "█"*50)
        logger.info("SECCIÓN 7: NET PPM")
        for d in all_dates:
            ok = download_daily_report(page, "netppm", d,
                 extra_filters={"Manufacturing": "Sourcing"}, account=account)
            if ok: total += 1
            random_delay()

    # ── 8. Catalog ────────────────────────────────────────────────────────────
    if should_run("catalog"):
        logger.info("\n" + "█"*50)
        logger.info("SECCIÓN 8: CATALOG")
        ok = download_static_report(page, "catalog",
             filters={"Manufacturing": "Sourcing"},
             h1_text="Catalog", account=account)
        if ok: total += 1

    logger.info(f"\n{'█'*50}")
    logger.info(f"✓ COMPLETO. Total archivos descargados: {total}")
    return page


def download_sales_range(page: Page, start: date, end: date):
    """Compatibilidad con código anterior."""
    download_all_reports(page, start, end)
